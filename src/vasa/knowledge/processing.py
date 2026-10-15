"""Bounded document ingestion; no filesystem extraction or external resources."""

import io
import re
import unicodedata
import zipfile
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 200_000
MAX_CHUNKS = 512
CHUNK_SIZE = 900
CHUNK_OVERLAP = 120


def normalize_text(text: str) -> str:
    if not isinstance(text, str) or len(text) > MAX_TEXT:
        raise ValueError("Knowledge text exceeds limit")
    if any(unicodedata.category(char) == "Cc" and char not in "\n\r\t" for char in text):
        raise ValueError("Knowledge text contains control characters")
    cleaned = re.sub(r"[ \t]+", " ", unicodedata.normalize("NFC", text))
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if not cleaned:
        raise ValueError("Knowledge text is empty")
    return cleaned


def process_text(text: str) -> list[str]:
    """Produce overlapping, character-bounded chunks in deterministic order."""
    text = normalize_text(text)
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = max(
                text.rfind("\n", start + CHUNK_SIZE // 2, end),
                text.rfind(" ", start + CHUNK_SIZE // 2, end),
            )
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if len(chunks) > MAX_CHUNKS:
            raise ValueError("Knowledge chunk limit exceeded")
        if end == len(text):
            break
        start = end - CHUNK_OVERLAP
    return chunks


def extract_document(name: str, data: bytes) -> str:
    """Accept strict UTF-8 TXT, nonencrypted PDF, or bounded DOCX only."""
    if not name or len(name) > 255 or "/" in name or "\\" in name:
        raise ValueError("Invalid document name")
    if not isinstance(data, bytes) or not data or len(data) > MAX_BYTES:
        raise ValueError("Document exceeds byte limit")
    suffix = PurePath(name).suffix.lower()
    if suffix == ".txt":
        text = data.decode("utf-8-sig", errors="strict")
    elif suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("Invalid PDF signature")
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted or len(reader.pages) > 200:
            raise ValueError("Encrypted or oversized PDF")
        parts: list[str] = []
        size = 0
        for page in reader.pages:
            # Guard oversized decompressed streams before parser extraction.
            contents = page.get_contents()
            if contents is not None and len(contents.get_data()) > 10 * 1024 * 1024:
                raise ValueError("PDF content stream exceeds limit")
            part = page.extract_text() or ""
            size += len(part)
            if size > MAX_TEXT:
                raise ValueError("PDF text exceeds limit")
            parts.append(part)
        text = "\n".join(parts)
    elif suffix == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 512 or sum(member.file_size for member in members) > 4 * MAX_BYTES:
                raise ValueError("DOCX archive exceeds limit")
            for member in members:
                path = PurePath(member.filename)
                if path.is_absolute() or ".." in path.parts or member.flag_bits & 1:
                    raise ValueError("Unsafe DOCX member")
                if member.file_size > 2 * MAX_BYTES:
                    raise ValueError("DOCX member exceeds limit")
            if "word/document.xml" not in archive.namelist():
                raise ValueError("Invalid DOCX structure")
            # Prevent XML entity declarations before handing content to native XML parser.
            for member in members:
                if member.filename.endswith((".xml", ".rels")):
                    payload = archive.read(member)
                    if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
                        raise ValueError("Unsafe DOCX XML")
        document = Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        parts.extend(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        text = "\n".join(parts)
    else:
        raise ValueError("Supported knowledge formats: TXT, PDF, DOCX")
    return normalize_text(text)
