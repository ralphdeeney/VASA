import io
import zipfile

import pytest
from docx import Document

from vasa.knowledge.processing import CHUNK_SIZE, extract_document, process_text


def test_chunks_preserve_bounded_overlap() -> None:
    text = " ".join(f"word{index}" for index in range(2000))
    chunks = process_text(text)
    assert len(chunks) > 5
    assert all(0 < len(chunk) <= CHUNK_SIZE for chunk in chunks)
    assert chunks[0].startswith("word0")
    assert chunks[-1].endswith("word1999")
    assert set(chunks[0].split()[-10:]) & set(chunks[1].split()[:30])


@pytest.mark.parametrize("text", ["", "  ", "a\x00b", "x" * 200001])
def test_invalid_text(text: str) -> None:
    with pytest.raises(ValueError):
        process_text(text)


def test_txt_unicode_and_strict_encoding() -> None:
    assert extract_document("notes.txt", b"\xef\xbb\xbfhello\nworld") == "hello\nworld"
    with pytest.raises(UnicodeDecodeError):
        extract_document("bad.txt", b"\xff")


@pytest.mark.parametrize(
    "name,data", [("../a.txt", b"a"), ("x.exe", b"a"), ("x.txt", b""), ("x.pdf", b"bad")]
)
def test_invalid_document(name: str, data: bytes) -> None:
    with pytest.raises(ValueError):
        extract_document(name, data)


def test_docx_paragraphs_tables() -> None:
    document = Document()
    document.add_paragraph("Booking policy")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Advance deposit required"
    stream = io.BytesIO()
    document.save(stream)
    assert extract_document("policy.docx", stream.getvalue()) == (
        "Booking policy\nAdvance deposit required"
    )


def test_docx_rejects_entities_and_traversal() -> None:
    for member, content in [
        ("../evil", b"x"),
        ("word/document.xml", b'<!DOCTYPE x [<!ENTITY x "boom">]>'),
    ]:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr(member, content)
        with pytest.raises(ValueError):
            extract_document("bad.docx", stream.getvalue())


def test_pdf_extracts_real_text_and_rejects_encrypted() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 700 Td (Business opening hours) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    assert extract_document("hours.pdf", output.getvalue()) == "Business opening hours"
    writer.encrypt("secret")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    with pytest.raises(ValueError, match="Encrypted"):
        extract_document("hours.pdf", encrypted.getvalue())
